"""由 run_all.py 的输出填写官方result模板并导出指定日期表格。需先运行 run_all.py 1..4。"""
import sys, json, pickle, os; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from core2 import *
OUT = '../../results/C'; os.makedirs(OUT, exist_ok=True)
d = load_data(); dates = d['dates']; p1 = np.tile(d['price1'], (365, 1)); p4 = d['price4']
o1 = pickle.load(open('out/1.pkl', 'rb')); o3 = pickle.load(open('out/3.pkl', 'rb'))
def fm(m): return f"{m//60 % 24}:{m%60:02d}" + ("+1" if m >= 1440 else "")
def lab(k): s = (k+1)*10; return f"{fm(s)}-{fm(s+10)}"     # 官方模板：数据时刻为区间起点
def blocks(a): return [float(a[i*24:(i+1)*24].sum()) for i in range(6)]
def emerg(em, thr=1e-6):
    res = []; k = 0
    while k < N:
        if em[k] > thr:
            j = k
            while j+1 < N and em[j+1] > thr: j += 1
            res.append((f"{fm((k+1)*10)}-{fm((j+1)*10+10)}", float(em[k:j+1].sum()))); k = j+1
        else: k += 1
    return res
HL = ('10:00', '12:00', '14:00', '16:00', '18:00', '20:00')

def write_year(res, name, price, q3):
    wb = openpyxl.load_workbook(f'{D}/{"result3" if q3 else "result2"}.xlsx')
    ws = wb['计划购电量']
    for o in res:
        if o['t'] < 31: continue
        r = o['t']-31+2; p = price[o['t']]
        for k in range(N): ws.cell(r, k+2, round(float(o['plan'][k]), 3))
        ws.cell(r, N+2, round(float(o['plan'].sum()), 3)); ws.cell(r, N+3, round(float((p*o['plan']).sum()), 3))
    if q3:
        ws = wb['调整购电量']
        for o in res:
            if o['t'] < 31: continue
            for k in range(36, N): ws.cell(o['t']-31+2, k+2, round(float(o['xf'][k]), 3))
    ws = wb['充放电量']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row: c.value = None
    r = 2
    for o in res:
        if o['t'] < 31: continue
        cb, db = blocks(o['c']), blocks(o['d']); ws.cell(r, 1, dates[o['t']].to_pydatetime())
        for i in range(6): ws.cell(r+i, 2, f"{4*i}:00-{4*i+4}:00"); ws.cell(r+i, 3, round(cb[i], 3)); ws.cell(r+i, 4, round(db[i], 3))
        ws.cell(r, 5, "0:00"); ws.cell(r, 6, round(float(o['E0']), 3)); ws.cell(r+1, 5, "24:00"); ws.cell(r+1, 6, round(float(o['Es'][-1]), 3)); r += 6
    ws = wb['紧急购电量']
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row: c.value = None
    r = 2
    for o in res:
        if o['t'] < 31: continue
        ws.cell(r, 1, dates[o['t']].to_pydatetime()); ps = emerg(o['em'])
        if not ps: r += 1; continue
        for s, a in ps: ws.cell(r, 2, s); ws.cell(r, 3, round(a, 3)); r += 1
    wb.save(f'{OUT}/{name}.xlsx')

def tables(res, price, q3):
    out = {}
    for ds in ('2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21'):
        t = int((dates == ds).nonzero()[0][0]); o = res[t]; p = price[t]
        out[ds] = dict(tab1={lab(k): float(o['plan'][k]) for k in range(N) if lab(k).split('-')[0] in HL},
                       tab1_final={lab(k): float(o['xf'][k]) for k in range(N) if lab(k).split('-')[0] in HL}, final_cost=float(o['cost_plan']), emerg_cost=float(o['cost_em']), plan_total=float(o['plan'].sum()), plan_cost=float((p*o['plan']).sum()), final_total=float(o['xf'].sum()),
                       charge=blocks(o['c']), discharge=blocks(o['d']), E0=float(o['E0']), E24=float(o['Es'][-1]),
                       emergency=emerg(o['em']), day_cost=float(o['cost_plan']+o['cost_em']))
    return out
S = {}
for name, res, price, q3 in (('result2', o1['Q2'], p1, False), ('result3', o1['Q3'], p1, True), ('result4-2', o3['Q4_2_fc'], p4, False), ('result4-3', o3['Q4_3_fc'], p4, True)):
    write_year(res, name, price, q3); S[name] = tables(res, price, q3)
# 问题1
A = json.load(open(f'{OUT}/analysis.json')); _, x, c, dd, E = solve_det(d['price1'], d['load1'], d['pv1'], 6000, 6000)
wb = openpyxl.load_workbook(f'{D}/result1.xlsx'); ws = wb['计划购电量']
for k in range(N): ws.cell(k+2, 2, round(float(x[k]), 4))
ws = wb['充放电量']; cb, db = blocks(c), blocks(dd)
for i in range(6): ws.cell(i+2, 2, round(cb[i], 4)); ws.cell(i+2, 3, round(db[i], 4))
ws.cell(2, 5, 6000); ws.cell(3, 5, round(float(E[-1]), 4)); wb.save(f'{OUT}/result1.xlsx')
json.dump(S, open(f'{OUT}/tables.json', 'w'), ensure_ascii=False, indent=1)
print('ok')
